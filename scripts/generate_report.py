"""Generate a Word (.docx) report summarizing the initial YOLOv8 training run and metrics."""
from pathlib import Path
from datetime import date

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parent.parent
RUN_DIR = ROOT / "runs" / "detect" / "runs" / "lettuce_health_v1"
OUT_PATH = ROOT / "Lettuce_Health_Initial_Training_Report.docx"

GREEN = RGBColor(0x1E, 0x7A, 0x34)
GRAY = RGBColor(0x55, 0x55, 0x55)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return table


def add_image_section(doc, title, img_path, caption, width_in=6.0):
    if not img_path.exists():
        return
    add_heading(doc, title, level=2)
    doc.add_picture(str(img_path), width=Inches(width_in))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = GRAY


def build():
    doc = Document()

    # Title
    title = doc.add_heading("Lettuce Health Detection — Initial Training Report", level=0)
    sub = doc.add_paragraph(f"YOLOv8 object detection baseline  |  Generated {date.today().isoformat()}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in sub.runs:
        r.font.color.rgb = GRAY
        r.font.size = Pt(11)

    # 1. Overview
    add_heading(doc, "1. Overview", level=1)
    doc.add_paragraph(
        "This report documents the first baseline training run for the lettuce health detection "
        "project. The goal of this stage was to validate the end-to-end pipeline — dataset "
        "preparation, model training, and evaluation — using the currently annotated images, "
        "classifying lettuce as either healthy or not healthy (leaf spot disease)."
    )

    # 2. Dataset summary
    add_heading(doc, "2. Dataset Summary", level=1)
    doc.add_paragraph(
        "Two separately annotated sources were merged into a single YOLO-format dataset with an "
        "80/20 train/validation split, stratified per class (seed=42):"
    )
    add_table(
        doc,
        ["Class", "Source images", "Train", "Val", "Total"],
        [
            ["healthy (0)", "naay_sakit_converted / WALAY SAKET", 34, 9, 43],
            ["not healthy (1)", "naaysaket_converted / NAAY SAKET", 111, 28, 139],
            ["Total", "", 145, 37, 182],
        ],
    )
    doc.add_paragraph()
    doc.add_paragraph(
        "Note: the dataset is small (182 images) and imbalanced (~76% not healthy vs. ~24% "
        "healthy). Results below should be read as an initial proof-of-concept rather than a "
        "production-ready benchmark."
    )

    # 3. Training configuration
    add_heading(doc, "3. Training Configuration", level=1)
    add_table(
        doc,
        ["Parameter", "Value"],
        [
            ["Model", "YOLOv8n (yolov8n.pt, pretrained)"],
            ["Image size", "640 x 640"],
            ["Batch size", 16],
            ["Epochs (max / actual)", "60 / 59 (early stopped, patience=20)"],
            ["Best epoch", "39"],
            ["Optimizer", "AdamW (auto-selected), lr0=0.001667, momentum=0.9"],
            ["Seed", 42],
            ["Hardware", "NVIDIA GeForce RTX 4050 Laptop GPU (6 GB VRAM)"],
            ["Software", "Ultralytics 8.4.184, torch 2.5.1+cu121 (CUDA)"],
            ["Training time", "~4.45 minutes (266.8 s)"],
        ],
    )

    # 4. Results
    add_heading(doc, "4. Results (Validation Set, Best Weights)", level=1)
    add_table(
        doc,
        ["Class", "Images", "Instances", "Precision", "Recall", "mAP50", "mAP50-95"],
        [
            ["healthy", 9, 9, 0.999, 1.000, 0.995, 0.974],
            ["not healthy", 28, 28, 0.919, 1.000, 0.988, 0.874],
            ["all (combined)", 37, 37, 0.959, 1.000, 0.991, 0.924],
        ],
    )
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Inference speed: ").bold = True
    p.add_run("0.3 ms preprocess + 7.7 ms inference + 1.9 ms postprocess per image on GPU "
               "(~100+ FPS capable).")

    add_heading(doc, "Interpretation", level=2)
    doc.add_paragraph(
        "The model reached very high precision/recall and mAP on the validation set, and a "
        "manual visual check of predictions confirmed the model correctly localizes the lettuce "
        "head and classifies health status in all sampled cases, rather than relying purely on "
        "background cues."
    )
    cav = doc.add_paragraph()
    cav.add_run("Caveats: ").bold = True
    cav.add_run(
        "the validation set is very small (37 images), so these metrics carry wide error "
        "margins — a couple of misclassifications would noticeably shift the scores. The two "
        "classes were also photographed in separate sessions, so there is some risk the model "
        "partly leverages incidental differences (lighting, background, framing) rather than "
        "disease symptoms alone. More images — especially additional healthy examples captured "
        "under varied conditions — are recommended before treating this as a reliable detector."
    )

    # 5. Visualizations
    add_heading(doc, "5. Visualizations", level=1)
    add_image_section(
        doc, "5.1 Training Curves", RUN_DIR / "results.png",
        "Box/class/DFL loss and precision/recall/mAP across training epochs."
    )
    add_image_section(
        doc, "5.2 Confusion Matrix", RUN_DIR / "confusion_matrix.png",
        "Predicted vs. true class counts on the validation set."
    )
    add_image_section(
        doc, "5.3 Sample Validation Predictions", RUN_DIR / "val_batch0_pred.jpg",
        "Model predictions on a batch of validation images (label + confidence shown per box)."
    )

    # 6. Artifacts / file locations
    add_heading(doc, "6. Run Artifacts", level=1)
    add_table(
        doc,
        ["Item", "Path"],
        [
            ["Best weights", "runs/detect/runs/lettuce_health_v1/weights/best.pt"],
            ["Last weights", "runs/detect/runs/lettuce_health_v1/weights/last.pt"],
            ["Full metrics (per epoch)", "runs/detect/runs/lettuce_health_v1/results.csv"],
            ["Merged dataset", "dataset/ (images/, labels/, data.yaml)"],
            ["Dataset prep script", "scripts/prepare_dataset.py"],
            ["Real-time detection script", "scripts/realtime_detect.py"],
        ],
    )

    # 7. Next steps
    add_heading(doc, "7. Next Steps", level=1)
    for item in [
        "Collect more images overall, and more 'healthy' examples specifically, to reduce class imbalance.",
        "Capture images under varied lighting/backgrounds to reduce the risk of the model learning incidental cues.",
        "Expand annotation to 3 classes: healthy, early leaf spot disease, and severe leaf spot disease.",
        "Re-evaluate on a larger, held-out validation set once more data is available.",
        "Use scripts/realtime_detect.py for live webcam/video testing of the current model.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(OUT_PATH)
    print(f"Saved report to {OUT_PATH}")


if __name__ == "__main__":
    build()
